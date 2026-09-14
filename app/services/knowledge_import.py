"""题库批量导入解析器：CSV / Excel（.xlsx）。

模板列（表头，支持常见别名）：
- 公司（company / 企业）
- 类别（category / 分类 / 板块 / 方向）
- 题目（question / 问题 / 试题）
- 标准答案（standard_answer / 答案 / 参考答案）
- 评分要点（scoring_points / 评分标准 / 要点 / 考点）—— 可选
"""
import csv
import io
from dataclasses import dataclass
from typing import Dict, List

from openpyxl import load_workbook


@dataclass
class ParsedItem:
    company: str
    category: str
    question: str
    standard_answer: str
    scoring_points: str = ""


# 标准字段 -> 可选别名（统一小写、去掉空格后匹配）
COLUMN_ALIASES: Dict[str, List[str]] = {
    "company": ["公司", "company", "企业", "单位"],
    "category": ["类别", "category", "分类", "板块", "方向", "领域"],
    "question": ["题目", "question", "问题", "试题", "题干"],
    "standard_answer": ["标准答案", "standardanswer", "答案", "参考答案", "标准解答"],
    "scoring_points": ["评分要点", "scoringpoints", "评分标准", "要点", "考点", "加分点"],
}

REQUIRED_FIELDS = ["company", "category", "question", "standard_answer"]


def _normalize(name: str) -> str:
    return name.strip().lower().replace(" ", "").replace("_", "")


def _build_column_map(fieldnames: List[str]) -> Dict[str, str]:
    """将表头映射为标准字段名；缺列时抛错说明。"""
    mapping: Dict[str, str] = {}
    for idx, raw in enumerate(fieldnames):
        key = _normalize(raw)
        for field, aliases in COLUMN_ALIASES.items():
            if key in [_normalize(a) for a in aliases]:
                mapping.setdefault(field, raw)  # 首个匹配优先
                break
    missing = [f for f in REQUIRED_FIELDS if f not in mapping]
    if missing:
        raise ValueError(
            f"模板缺少必要列：{', '.join(missing)}。"
            "请使用模板列：公司、类别、题目、标准答案、评分要点(可选)"
        )
    return mapping


def rows_to_items(rows: List[Dict[str, str]]) -> List[ParsedItem]:
    """把 [ {列名: 值}, ... ] 转为 ParsedItem 列表，跳过空行。"""
    if not rows:
        return []
    mapping = _build_column_map(list(rows[0].keys()))
    items: List[ParsedItem] = []
    for row in rows:
        def get(field: str) -> str:
            col = mapping.get(field)
            return (row.get(col) or "").strip() if col else ""

        company = get("company")
        category = get("category")
        question = get("question")
        standard_answer = get("standard_answer")
        scoring_points = get("scoring_points")
        if not question or not standard_answer:
            continue  # 跳过不完整的空行
        items.append(
            ParsedItem(
                company=company or "未标注公司",
                category=category or "未分类",
                question=question,
                standard_answer=standard_answer,
                scoring_points=scoring_points,
            )
        )
    return items


def _decode(content: bytes) -> str:
    """尝试多种编码（优先 UTF-8 BOM，兼容中文 Excel 的 GBK）。"""
    for enc in ("utf-8-sig", "gbk", "utf-8"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ValueError("无法识别文件编码，请用 UTF-8 或 GBK 保存")


def parse_csv(content: bytes) -> List[ParsedItem]:
    text = _decode(content)
    reader = csv.DictReader(io.StringIO(text))
    rows: List[Dict[str, str]] = []
    for row in reader:
        if row is None:
            continue
        # DictReader 的值为 None 时归一为空串
        rows.append({k: (v or "") for k, v in row.items()})
    return rows_to_items(rows)


def parse_excel(content: bytes) -> List[ParsedItem]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    raw = list(ws.iter_rows(values_only=True))
    wb.close()
    if not raw:
        return []
    headers = [(str(c).strip() if c is not None else "") for c in raw[0]]
    rows: List[Dict[str, str]] = []
    for r in raw[1:]:
        values = [(str(c).strip() if c is not None else "") for c in r]
        # 补齐长度，避免列数不一致
        while len(values) < len(headers):
            values.append("")
        rows.append(dict(zip(headers, values)))
    return rows_to_items(rows)


def parse_upload(filename: str, content: bytes) -> List[ParsedItem]:
    """按扩展名分发解析；只接受 .csv 与 .xlsx。"""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "csv":
        return parse_csv(content)
    if ext == "xlsx":
        return parse_excel(content)
    raise ValueError("仅支持 .csv 或 .xlsx 文件，请先下载模板")
